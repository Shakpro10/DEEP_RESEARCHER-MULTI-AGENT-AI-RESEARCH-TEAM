# ============================================
# STEP 1: import necessary modules and classes
# ============================================
import re
import tempfile
import uuid
from pathlib import Path
import gradio as gr
from docx import Document
from fpdf import FPDF

# Define the output directory for generated reports
OUTPUT_DIR = Path(tempfile.gettempdir()) / "pulse_reports"
OUTPUT_DIR.mkdir(exist_ok=True)

# Define regex patterns for headings, bold, and italic markdown
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)")
_BOLD_RE = re.compile(r"\*\*(.*?)\*\*")
_ITALIC_RE = re.compile(r"\*(.*?)\*")

# Define a helper function to strip inline markdown formatting from text
def _strip_inline_markdown(text: str) -> str:
    text = _BOLD_RE.sub(r"\1", text)
    text = _ITALIC_RE.sub(r"\1", text)
    return text

# ==============================================================
# STEP 2: Define functions to build DOCX file from markdown text
# ==============================================================
def build_docx(markdown_text: str, path: str) -> None:
    """Very lightweight markdown -> docx converter (headings + bullets + text)."""
    doc = Document()
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph("")
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(_strip_inline_markdown(heading.group(2)), level=level)
            continue

        if line.startswith(("- ", "* ")):
            doc.add_paragraph(_strip_inline_markdown(line[2:]), style="List Bullet")
            continue

        doc.add_paragraph(_strip_inline_markdown(line))

    doc.save(path)


# =============================================================
# STEP 3: Define functions to build PDF file from markdown text
# =============================================================
def build_pdf(markdown_text: str, path: str) -> None:
    """Very lightweight markdown -> PDF converter (headings + bullets + text).

    fpdf2's core fonts are latin-1 only, so any characters outside that range
    (smart quotes, emoji, non-Latin scripts, etc.) are safely replaced rather
    than raising -- acceptable for a report export, not meant for perfect
    typographic fidelity.
    """
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=11)

    heading_sizes = {1: 18, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}

    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        pdf.set_x(pdf.l_margin)  # multi_cell can leave the cursor at the right
        # margin; always reset before writing the next line.

        if not line:
            pdf.ln(4)
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            size = heading_sizes.get(len(heading.group(1)), 11)
            pdf.set_font("Helvetica", "B", size)
            safe = _strip_inline_markdown(heading.group(2)).encode(
                "latin-1", "replace"
            ).decode("latin-1")
            pdf.multi_cell(0, 8, safe)
            pdf.set_font("Helvetica", size=11)
            continue

        if line.startswith(("- ", "* ")):
            safe = _strip_inline_markdown(line[2:]).encode(
                "latin-1", "replace"
            ).decode("latin-1")
            pdf.multi_cell(0, 6, f"-  {safe}")
            continue

        safe = _strip_inline_markdown(line).encode("latin-1", "replace").decode(
            "latin-1"
        )
        pdf.multi_cell(0, 6, safe)

    pdf.output(path)


# =====================================================================
# STEP 4: Define a function to generate downloadable PDF and DOCX files
# =====================================================================
def generate_downloads(report_text: str):
    """Gradio callback: builds fresh PDF/DOCX files and reveals the two
    download buttons. Returns gr.update() pairs for (pdf_download, docx_download).
    """
    if not report_text or not report_text.strip():
        return gr.update(value=None, visible=False), gr.update(value=None, visible=False)

    stamp = uuid.uuid4().hex[:8]
    pdf_path = OUTPUT_DIR / f"research_report_{stamp}.pdf"
    docx_path = OUTPUT_DIR / f"research_report_{stamp}.docx"

    build_pdf(report_text, str(pdf_path))
    build_docx(report_text, str(docx_path))

    return (
        gr.update(value=str(pdf_path), visible=True),
        gr.update(value=str(docx_path), visible=True),
    )
